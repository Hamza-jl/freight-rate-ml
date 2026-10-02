"""Build reports/Freight_Rate_Report.docx from the metrics and figures.

    python build_report.py   (after train.py, predict.py, score.py, eda.py)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent / "src"))
from freight_rate import config  # noqa: E402

CHART = config.ROOT / "scorer_results" / "candidate_december.png"


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    doc.add_paragraph()


def figure(doc, path, caption, width=6.2):
    doc.add_picture(str(path), width=Inches(width))
    p = doc.add_paragraph(caption)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)


def main():
    m = json.loads(config.METRICS_PATH.read_text())
    bt = pd.read_csv(config.REPORT_DIR / "backtest.csv")
    mean = m["backtest_mean"]
    final = mean["LightGBM, final features"]
    base = mean["Baseline: rate/mile by equipment x distance band"]
    naive = mean["LightGBM, raw features"]
    fig = config.FIGURE_DIR

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Freight Rate Prediction: Validation Approach & Results", 0)
    doc.add_paragraph(
        f"Task: predict posted_rate for 12,000 loads dated Nov–Dec 2025, using {m['n_train_rows']:,} labelled loads from "
        f"Jan–Oct 2025. Final model: LightGBM on log(rate). Rolling time-based backtest: mean MAPE {final['MAPE']:.2f}% "
        f"(MAE ${final['MAE']:.0f}), against {base['MAPE']:.2f}% for a rate-per-mile baseline."
    )

    doc.add_heading("1. Data exploration & quality", 1)
    doc.add_paragraph(
        "Distance drives price. Rate per mile falls from about $2.80 under 200 miles to about $1.90 over 2,000 miles. "
        "Reefer costs about 13% more per mile than Dry Van, and Flatbed about 8% more. Heavier loads price slightly higher. "
        "Once these are accounted for, clean labels sit within ±4% of a simple model for 90% of loads, so the noise is low "
        "and the real difficulty lies in the data-quality issues and the market-level drift over time."
    )
    table(
        doc,
        ["Issue", "Rows affected", "Handling"],
        [
            ["posted_rate corrupted (0.2–0.45× or 2–5× the expected value)", f"{m['n_corrupted_labels']} (1.4%)",
             "Flagged when >1.5× away from a 5-fold out-of-fold model; dropped from training and evaluation"],
            ["Negative weight", "292 train / 145 validation", "Sign error → absolute value (these loads price the same as positive-weight loads)"],
            ["Missing weight", "300 / 165", "Training median weight for the equipment type"],
            ["Missing market_index", "374 / 249", "That day's median (98% of its variance is explained by date)"],
            ["Validation cities never seen in training", "8 cities", "No city-ID features; the model uses coordinates and distance"],
            ["Coordinates clipped at lat 25.5 / lon −69.5", "3 cities", "Consistent per city, so kept as relative geography"],
        ],
    )
    figure(doc, fig / "corrupted_labels.png", "Figure 1: Corrupted labels form two separate bands, 2–5× above and below the clean diagonal.", 5.2)
    doc.add_paragraph(
        "market_index is a trap feature. 98% of its variance is a daily, market-wide value: a 7-day cycle on top of a "
        "spring/summer hump. That level moves about 40% across the year while rates move about 10%, and the same index "
        "value lines up with different rate levels in different months. A tree model fed the raw value learns splits "
        "tied to specific months that do not carry over to future months. It is therefore split into a daily level, "
        "which is dropped, and each load's deviation from that level, which is kept. quote_signal is decomposed the same way."
    )
    figure(doc, fig / "market_vs_rates.png", "Figure 2: Monthly medians indexed to January. market_index is not a reliable price-level signal.")
    figure(doc, fig / "daily_market_index.png", "Figure 3: The daily market_index has a strong weekly cycle; validation covers a calm regime.")

    doc.add_heading("2. Train / validation split approach", 1)
    doc.add_paragraph(
        "The data is time-ordered, and the real task is to forecast the next two months (train Jan–Oct, predict Nov–Dec). "
        "A random split would let the model see every month's market regime during training, which leaks future "
        "information. Validation therefore mirrors deployment with a rolling-origin backtest. Each fold trains on all "
        "data up to a cut-off month and tests on the following two months:"
    )
    rows = []
    for fold in dict.fromkeys(bt["fold"]):
        sub = bt[bt["fold"] == fold].set_index("model")
        rows.append([fold, f"{sub['n'].iloc[0]:,}"] + [f"{sub.loc[k, 'MAPE']:.2f}%" for k in sub.index])
    rows.append(["Mean", "", f"{base['MAPE']:.2f}%", f"{naive['MAPE']:.2f}%", f"{final['MAPE']:.2f}%"])
    table(doc, ["Fold", "Test rows", "Baseline", "LightGBM raw", "LightGBM final"], rows)
    doc.add_paragraph(
        f"Choices and checks: (1) corrupted labels are removed from the test folds too, so the metrics measure "
        f"real pricing error. Scored against the raw labels, the Sep–Oct holdout gives MAPE "
        f"{m['holdout_sep_oct_raw_labels']['MAPE']:.2f}%, which is driven entirely by the 2–5× bad rows. "
        f"(2) A random 80/20 split gives {m['random_split']['MAPE']:.2f}% MAPE, lower than the time-based "
        f"{final['MAPE']:.2f}%, which confirms that a random split would be over-optimistic. "
        f"(3) Day-level market statistics are computed from the features of all loads on the same date, never from "
        f"labels. That information is also available when the validation set is scored. "
        f"(4) Features, hyperparameters and the outlier threshold were all chosen on the backtest. The final model "
        f"is then refit on all clean Jan–Oct data."
    )
    figure(doc, fig / "backtest_mape.png", "Figure 4: Backtest MAPE per fold. The final features are the only variant that beats the baseline in every fold.")

    doc.add_heading("3. Model", 1)
    doc.add_paragraph(
        "LightGBM gradient-boosted trees on log(posted_rate), with 1,500 trees, 15 leaves, learning rate 0.03 and row "
        "and column subsampling of 0.8. Reasons: the rate-per-mile curve is non-linear in distance and interacts with "
        "equipment and geography, which trees capture without hand-built terms. The log target turns errors into "
        "relative errors, which suits rates from about $60 to $8k. A Ridge model on the log target scored 3.7% "
        "MAPE in the backtest. Hyperparameter variations (leaves 15–63, L1/Huber objectives, recency weighting) all fell "
        "within ±0.1 percentage points of MAPE, so the simpler configuration was kept."
    )
    figure(doc, fig / "feature_importance.png", "Figure 5: Distance dominates; geography and quote_signal refine it.", 5.2)

    doc.add_heading("4. December prediction chart (fixed inputs)", 1)
    figure(doc, CHART, "Figure 6: Output of the provided score.py. Lexington → Fort Wayne, 360 mi, Dry Van, 32,000 lb.")
    doc.add_paragraph(
        "The chart file provides only lane, distance, equipment, weight and date. Coordinates come from the city lookup. "
        "market_index and quote_signal are set to that day's median across all validation loads, so the load-level "
        "deviations are 0 and the chart shows the model's expected rate for this fixed load as the market moves. "
        "Predictions stay between about $792 and $816, in line with the 21 Dry Van loads on this lane in training (median $808, range $758–$934)."
        "The weekly pattern (lower at weekends, highest mid-week) matches the day-of-week effect in the training data. "
        "The step on 31 December comes from that day's unusually low market quote level, which the model associates "
        "with October's pricing regime."
    )

    doc.add_heading("5. Limitations & next steps", 1)
    for text in [
        "The training data has no November or December history, so holiday or year-end seasonality cannot be learned. A second year of data, or an explicit seasonal prior, would address this.",
        "The market level over the next two months is the dominant source of error (fold MAPE ranges from 1.6% to 2.6%). A dedicated time-series model of the daily market level would be the next step.",
        "The corruption flag is a fixed threshold. In production it should feed a data-quality alert upstream rather than silently dropping rows.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    out = config.REPORT_DIR / "Freight_Rate_Report.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
